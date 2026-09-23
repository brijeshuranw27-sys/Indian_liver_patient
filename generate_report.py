from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.shared import Inches

ROOT = Path(__file__).parent
DATA = ROOT / "liver_patient.csv"
IMAGE_DIR = ROOT / "report_images"


def generate() -> None:
    """Generate a short DOCX report and supporting chart images.

    Creates report_images/ and prediction_report.docx in the project root.
    """
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)

    if not DATA.exists():
        raise FileNotFoundError(f"Dataset not found at {DATA}. Place the CSV in the project root.")

    data = pd.read_csv(DATA)

    # Create a simple class distribution chart
    ax = data.groupby("Dataset")["Age"].count().plot(kind="bar", color=["#dd745d", "#4c8c79"])
    ax.set_title("Example records by target class")
    ax.set_xlabel("Dataset class")
    ax.set_ylabel("Records")
    chart_path = IMAGE_DIR / "class_distribution.png"
    ax.figure.tight_layout()
    ax.figure.savefig(chart_path, dpi=160)
    plt.close(ax.figure)

    # Build the Word document
    document = Document()
    document.add_heading("Liver Patient Prediction Project", 0)
    document.add_paragraph(
        "A reproducible machine-learning prototype for exploring liver laboratory profiles."
    )
    document.add_heading("Purpose", level=1)
    document.add_paragraph(
        "This project demonstrates data preparation, model training, and a Flask prediction workflow."
        " It is intended for educational purposes only and not for clinical diagnosis."
    )
    document.add_heading("Data and model", level=1)
    document.add_paragraph(
        f"The example dataset contains {len(data)} records and {len(data.columns) - 1} input measurements. "
        "Missing values are imputed with the median; a class-balanced Random Forest classifier is used for prediction."
    )

    # Insert chart
    document.add_picture(str(chart_path), width=Inches(5.8))

    document.add_heading("Reproduction", level=1)
    document.add_paragraph(
        "Reproduce locally: create a virtual environment, install requirements, run 'python train_model.py' to produce the model artifact,"
        " then run 'python app.py' to start the prediction service."
    )

    out_path = ROOT / "prediction_report.docx"
    document.save(out_path)
    print(f"Generated {out_path} and {chart_path}")


if __name__ == "__main__":
    generate()
